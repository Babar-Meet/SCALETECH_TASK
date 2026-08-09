import os
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
from config import config

def export_links_to_excel(links):
    if not os.path.exists(config.REPORTS_DIRECTORY):
        os.makedirs(config.REPORTS_DIRECTORY, exist_ok=True)
        
    filepath = os.path.join(config.REPORTS_DIRECTORY, "imported_links.xlsx")
    
    wb = Workbook()
    ws = wb.active
    ws.title = "WhatsApp Links"
    
    headers = [
        "Date", "Sender", "Message", "URL", "Title", "Channel", 
        "User Category", "AI Category", "AI Confidence", "Language", 
        "Status", "Local Path"
    ]
    ws.append(headers)
    
    # Style header
    for col in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col)
        cell.font = Font(bold=True)
        
    for link in links:
        ws.append([
            link.message_date.strftime("%Y-%m-%d %H:%M:%S") if link.message_date else "",
            link.sender,
            link.original_message,
            link.extracted_url,
            link.audio_title or "",
            link.channel_name or "",
            link.category or "",
            link.ai_category or "",
            link.ai_confidence or "",
            link.ai_language or "",
            link.status,
            link.audio_path or ""
        ])
        
    # Wrap text for Message column (C)
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=3, max_col=3):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True)
            
    # Set column widths
    ws.column_dimensions['A'].width = 20
    ws.column_dimensions['B'].width = 20
    ws.column_dimensions['C'].width = 40
    ws.column_dimensions['D'].width = 30
    ws.column_dimensions['E'].width = 30
    ws.column_dimensions['F'].width = 20
    ws.column_dimensions['G'].width = 15
    ws.column_dimensions['H'].width = 15
    ws.column_dimensions['I'].width = 15
    ws.column_dimensions['J'].width = 15
    ws.column_dimensions['K'].width = 15
    ws.column_dimensions['L'].width = 50
    
    # Freeze header
    ws.freeze_panes = "A2"
    
    wb.save(filepath)
    return filepath
